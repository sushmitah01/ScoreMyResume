import pytest
import docx
from fpdf import FPDF
from backend.services.parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_resume,
)


@pytest.fixture
def sample_pdf(tmp_path):
    pdf_path = tmp_path / "sample_resume.pdf"
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="John Doe - Python Developer with 5 years of experience.")
    pdf.output(str(pdf_path))
    return str(pdf_path)


@pytest.fixture
def sample_docx(tmp_path):
    docx_path = tmp_path / "sample_resume.docx"
    document = docx.Document()
    document.add_paragraph("Jane Smith - Backend Engineer skilled in Docker and AWS.")
    document.save(str(docx_path))
    return str(docx_path)


def test_extract_text_from_pdf_returns_readable_text(sample_pdf):
    text = extract_text_from_pdf(sample_pdf)
    assert "Python Developer" in text


def test_extract_text_from_docx_returns_readable_text(sample_docx):
    text = extract_text_from_docx(sample_docx)
    assert "Backend Engineer" in text


def test_parse_resume_routes_pdf_correctly(sample_pdf):
    text = parse_resume(sample_pdf)
    assert "Python Developer" in text


def test_parse_resume_routes_docx_correctly(sample_docx):
    text = parse_resume(sample_docx)
    assert "Backend Engineer" in text


def test_parse_resume_raises_error_for_unsupported_extension(tmp_path):
    fake_file = tmp_path / "resume.txt"
    fake_file.write_text("plain text resume")

    with pytest.raises(ValueError):
        parse_resume(str(fake_file))